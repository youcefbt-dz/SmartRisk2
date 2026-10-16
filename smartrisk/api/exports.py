"""تصدير القوائم المالية: CSV وExcel."""
from __future__ import annotations

import csv
import io

from flask import Blueprint, Response

from .. import repo
from ..auth import current_user, login_required
from ..core import registry
from . import bad_request

bp = Blueprint("api_exports", __name__, url_prefix="/api/projects")

SECTIONS = {"income": "قائمة الدخل", "cashflow": "التدفقات النقدية", "balance": "الميزانية"}


def _statements(pid: str):
    proj = repo.get_project(current_user()["id"], pid)  # type: ignore[index]
    if not proj:
        return None, None
    return proj, registry.run("statements", proj["data"], {})


def _safe_cell(v):
    """يمنع حقن الصيغ في Excel/CSV: النص الذي يبدأ برموز الصيغ يُسبق بفاصلة عليا."""
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + v
    return v


@bp.get("/<pid>/export/<section>.csv")
@login_required
def export_csv(pid, section):
    if section not in SECTIONS:
        return bad_request("قسم غير معروف.", 404)
    proj, st = _statements(pid)
    if not proj:
        return bad_request("المشروع غير موجود.", 404)
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["البند"] + [f"ش{m}" for m in st["months"]])
    for line in st[section]:
        w.writerow([_safe_cell(line["label"])] + [round(v, 2) for v in line["monthly"]])
    data = "﻿" + buf.getvalue()  # BOM لعرض العربية صحيحاً في Excel
    return Response(data, mimetype="text/csv; charset=utf-8", headers={
        "Content-Disposition": f"attachment; filename=smartrisk-{section}.csv"})


@bp.get("/<pid>/export.xlsx")
@login_required
def export_xlsx(pid):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        return bad_request("تصدير Excel يتطلب مكتبة openpyxl.", 501)
    proj, st = _statements(pid)
    if not proj:
        return bad_request("المشروع غير موجود.", 404)
    wb = Workbook()
    wb.remove(wb.active)
    head_fill = PatternFill("solid", fgColor="231D70")
    for key, title in SECTIONS.items():
        ws = wb.create_sheet(title)
        ws.sheet_view.rightToLeft = True
        ws.append(["البند"] + [f"السنة {y}" for y in st["years"]] + [f"ش{m}" for m in st["months"]])
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = head_fill
            c.alignment = Alignment(horizontal="center")
        for line in st[key]:
            row = [_safe_cell(line["label"])] + [round(v, 2) for v in line["annual"]] + [round(v, 2) for v in line["monthly"]]
            ws.append(row)
            if line["total"]:
                for c in ws[ws.max_row]:
                    c.font = Font(bold=True)
        ws.column_dimensions["A"].width = 44
        ws.freeze_panes = "B2"
        for row in ws.iter_rows(min_row=2, min_col=2):
            for c in row:
                c.number_format = "#,##0"
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename=smartrisk-{pid}.xlsx"})
